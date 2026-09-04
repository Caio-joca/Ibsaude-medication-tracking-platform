from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from extensions import db
from models import Distribuicao, Medicamento, Movimentacao, UnidadeDestino, Usuario
from utils.seguranca import perfis_permitidos


relatorios_bp = Blueprint("relatorios", __name__, url_prefix="/relatorios")


def movimentos_filtrados():
    consulta = Movimentacao.query
    inicio = request.args.get("inicio")
    fim = request.args.get("fim")
    medicamento = request.args.get("medicamento", type=int)
    lote = request.args.get("lote", "").strip()
    unidade = request.args.get("unidade", type=int)
    usuario = request.args.get("usuario", type=int)
    try:
        if inicio:
            consulta = consulta.filter(Movimentacao.dt_movimentacao >= date.fromisoformat(inicio))
        if fim:
            limite_fim = datetime.combine(date.fromisoformat(fim) + timedelta(days=1), time.min)
            consulta = consulta.filter(Movimentacao.dt_movimentacao < limite_fim)
    except ValueError:
        pass
    if medicamento:
        consulta = consulta.filter(Movimentacao.id_medicamento == medicamento)
    if lote:
        consulta = consulta.join(Movimentacao.estoque).filter_by(lote=lote)
    if unidade:
        consulta = consulta.join(Movimentacao.distribuicao).filter(Distribuicao.id_destino == unidade)
    if usuario:
        consulta = consulta.filter(Movimentacao.id_usuario == usuario)
    return consulta.order_by(Movimentacao.dt_movimentacao.desc()).all()


def linha_movimento(movimento):
    unidade = movimento.distribuicao.unidade.nome if movimento.distribuicao else "-"
    lote = movimento.estoque.lote if movimento.estoque else "-"
    return [
        movimento.dt_movimentacao.strftime("%d/%m/%Y %H:%M"),
        movimento.tipo_movimentacao,
        movimento.medicamento.nome,
        lote,
        movimento.quantidade,
        unidade,
        movimento.usuario.nome if movimento.usuario else "-",
        float(movimento.vlr_movimentacao),
    ]


@relatorios_bp.get("")
@perfis_permitidos("A", "G", "U")
def index():
    movimentos = movimentos_filtrados()
    total = sum(float(item.vlr_movimentacao) for item in movimentos)
    entradas = [m for m in movimentos if m.tipo_movimentacao == "E"]
    saidas = [m for m in movimentos if m.tipo_movimentacao == "S"]
    por_unidade = {}
    for movimento in saidas:
        nome = movimento.distribuicao.unidade.nome if movimento.distribuicao else "Sem unidade"
        resumo = por_unidade.setdefault(nome, {"quantidade": 0, "valor": 0})
        resumo["quantidade"] += movimento.quantidade
        resumo["valor"] += float(movimento.vlr_movimentacao)
    return render_template(
        "relatorios/index.html",
        movimentos=movimentos,
        total=total,
        medicamentos=Medicamento.query.order_by(Medicamento.nome).all(),
        unidades=UnidadeDestino.query.order_by(UnidadeDestino.nome).all(),
        usuarios=Usuario.query.order_by(Usuario.nome).all(),
        entradas_qtd=sum(item.quantidade for item in entradas),
        entradas_valor=sum(float(item.vlr_movimentacao) for item in entradas),
        saidas_qtd=sum(item.quantidade for item in saidas),
        saidas_valor=sum(float(item.vlr_movimentacao) for item in saidas),
        por_unidade=sorted(por_unidade.items()),
        query_string=request.query_string.decode("utf-8"),
    )


@relatorios_bp.get("/movimentacoes.pdf")
@perfis_permitidos("A", "G", "U")
def pdf():
    dados = [linha_movimento(m) for m in movimentos_filtrados()]
    fluxo = BytesIO()
    documento = SimpleDocTemplate(fluxo, pagesize=landscape(A4), rightMargin=1 * cm, leftMargin=1 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
    estilos = getSampleStyleSheet()
    elementos = [Paragraph("IBSAÚDE - Relatório de movimentações", estilos["Title"]), Spacer(1, 0.4 * cm)]
    tabela = Table([["Data", "Tipo", "Medicamento", "Lote", "Qtd.", "Unidade", "Responsável", "Valor (R$)"]] + dados, repeatRows=1, colWidths=[2.8*cm, 1*cm, 4.2*cm, 2.4*cm, 1.2*cm, 3.4*cm, 3.4*cm, 2.2*cm])
    tabela.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0d6efd")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), 0.4, colors.grey), ("FONTSIZE", (0,0), (-1,-1), 8), ("VALIGN", (0,0), (-1,-1), "TOP"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f2f6fa")])]))
    elementos.append(tabela)
    documento.build(elementos)
    fluxo.seek(0)
    return send_file(fluxo, mimetype="application/pdf", as_attachment=True, download_name="relatorio_movimentacoes.pdf")


@relatorios_bp.get("/movimentacoes.xlsx")
@perfis_permitidos("A", "G", "U")
def excel():
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "Movimentações"
    cabecalho = ["Data", "Tipo", "Medicamento", "Lote", "Quantidade", "Unidade", "Responsável", "Valor (R$)"]
    planilha.append(cabecalho)
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D6EFD")
    for movimento in movimentos_filtrados():
        planilha.append(linha_movimento(movimento))
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions
    larguras = [20, 8, 32, 18, 12, 28, 28, 16]
    for indice, largura in enumerate(larguras, 1):
        planilha.column_dimensions[chr(64 + indice)].width = largura
    for celula in planilha["H"][1:]:
        celula.number_format = 'R$ #,##0.00'
    fluxo = BytesIO()
    pasta.save(fluxo)
    fluxo.seek(0)
    return send_file(fluxo, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="relatorio_movimentacoes.xlsx")
